#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将日报 Markdown 转换为 Word（.docx）。

用法：
    py -3 references/md_to_docx.py reports/20261002.md
    py -3 references/md_to_docx.py reports/20261002.md reports/20261002.docx

依赖 python-docx（仅安装在 Windows Python 3.13，需用 py -3 运行）。
仅解析日报模板的固定结构：# / ## / ### 标题、- 列表项、普通段落。
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

CJK_FONT = "微软雅黑"


def set_cjk_font(doc):
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = "Calibri"
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), CJK_FONT)
    doc.styles["Normal"].font.size = Pt(11)


def clean(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).strip()


def convert(md_path, docx_path):
    doc = Document()
    set_cjk_font(doc)

    for raw in Path(md_path).read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(clean(line[4:]), level=2)
        elif line.startswith("## "):
            doc.add_heading(clean(line[3:]), level=1)
        elif line.startswith("# "):
            doc.add_heading(clean(line[2:]), level=0)
        elif line.startswith("- "):
            doc.add_paragraph(clean(line[2:]), style="List Bullet")
        else:
            doc.add_paragraph(clean(line))

    doc.save(docx_path)
    return docx_path


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except AttributeError:
            pass
    if len(sys.argv) < 2:
        sys.exit("用法: py -3 references/md_to_docx.py <日报.md> [输出.docx]")
    md_path = sys.argv[1]
    docx_path = sys.argv[2] if len(sys.argv) > 2 else str(Path(md_path).with_suffix(".docx"))
    convert(md_path, docx_path)
    print(f"Word 日报已生成 -> {docx_path}")


if __name__ == "__main__":
    main()
